"""Convert six airline datasets to linked CSV files. Python 3.10+.

PyCharm: open this file and press Run. Edit INPUT_DIR / OUTPUT_ROOT below.
Only YourBoardingPassDotAero.zip is opened; no PDF or Airlines-All.zip.
No people are merged or assigned global user IDs here.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import time
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from contextlib import ExitStack
from datetime import date, datetime, time as datetime_time
from pathlib import Path
import ijson
import simplejson
import yaml
from openpyxl import load_workbook

INPUT_DIR = Path(__file__).resolve().parent / "input_dir"
OUTPUT_ROOT = Path(__file__).resolve().parent / "converted"
SOURCES = {
    "boarding": "BoardingData.csv",
    "sirena": "Sirena-export-fixed.tab",
    "pointz": "PointzAggregator-AirlinesData.xml",
    "forum": "FrequentFlyerForum-Profiles.json",
    "skyteam": "SkyTeam-Exchange.yaml",
    "boardingpass": "YourBoardingPassDotAero.zip",
}
VERSION = "1.1"
DATA_COLUMNS = [
    "observation_id", "record_kind", "first_name", "middle_name", "last_name",
    "full_name", "birth_date_raw", "birth_date_iso", "sex", "passport_raw",
    "passport_key", "loyalty_cards_json", "flight_date", "flight_time",
    "flight_number", "departure_airport", "departure_city", "departure_country",
    "arrival_airport", "arrival_city", "arrival_country", "arrival_date",
    "arrival_time", "booking_code", "ticket_number", "codeshare_raw",
    "travel_class", "fare", "baggage_status", "baggage_allowance", "seat",
    "meal", "agent", "flight_status", "gate", "boarding_sequence", "honorific", "issues_json",
]
RECORD_COLUMNS = ["source_record_id", "source_file", "source_position", "payload_format", "payload"]
LINK_COLUMNS = ["observation_id", "source_record_id", "source_subject_id", "source_path", "loyalty_scope"]
SIRENA_COLUMNS = [
    ("PaxName",0,60),("PaxBirthDate",60,72),("DepartDate",72,84),
    ("DepartTime",84,96),("ArrivalDate",96,108),("ArrivalTime",108,120),
    ("Flight",120,126),("CodeSh",126,132),("From",132,138),("Dest",138,144),
    ("Code",144,150),("e-Ticket",150,168),("TravelDoc",168,180),
    ("Seat",180,186),("Meal",186,192),("TrvCls",192,198),
    ("Fare",198,204),("Baggage",204,216),("PaxAdditionalInfo",216,276),
    ("AgentInfo",276,336),
]
BOARDING_COLUMNS = (
    "PassengerFirstName;PassengerSecondName;PassengerLastName;PassengerSex;"
    "PassengerBirthDate;PassengerDocument;BookingCode;TicketNumber;Baggage;"
    "FlightDate;FlightTime;FlightNumber;CodeShare;Destination"
).split(";")
MISSING_ID = {"", "n/a", "not presented", "null", "none"}


def compact(value) -> str:
    # ijson returns Decimal for fractional JSON numbers; retain numeric precision.
    return simplejson.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False, use_decimal=True)


def text(value) -> str:
    return "" if value is None else str(value)


def passport_key(value) -> str:
    """Conservative formatting normalization; never pad/truncate or cast to int."""
    value = unicodedata.normalize("NFKC", text(value)).strip()
    if value.casefold() in MISSING_ID:
        return ""
    result = re.sub(r"[\s\-]", "", value).upper()
    return result if re.fullmatch(r"[A-ZА-ЯЁ0-9]+", result) else ""


def card(program, number, **extra) -> dict:
    p, n = text(program), text(number)
    pk = unicodedata.normalize("NFKC", p).strip().upper()
    nk = re.sub(r"\s", "", unicodedata.normalize("NFKC", n))
    valid = bool(re.fullmatch(r"[A-Z0-9]{2,8}", pk) and re.fullmatch(r"\d+", nk))
    return {"program_raw": p, "number_raw": n, "program_key": pk if valid else "",
            "number_key": nk if valid else "", "card_key": f"{pk}:{nk}" if valid else "", **extra}


def combined_card(value, **extra) -> dict:
    match = re.fullmatch(r"\s*([A-Za-z0-9]+)\s+(\d+)\s*", text(value))
    if not match:
        return card("", "", combined_raw=value, **extra)
    return card(match[1], match[2], combined_raw=value, **extra)


def iso_birth(value, fmt, issues) -> str:
    value = text(value)
    if value.strip().casefold() in MISSING_ID:
        return ""
    try:
        return datetime.strptime(value.strip(), fmt).date().isoformat()
    except ValueError:
        issues.append("birth_date_not_parsed")
        return ""


def person(first="", middle="", last="", full="", birth="", birth_fmt="%Y-%m-%d", sex="", passport="", cards=None):
    issues = []
    p = dict(first_name=text(first), middle_name=text(middle), last_name=text(last),
             full_name=text(full), birth_date_raw=text(birth), sex=text(sex),
             passport_raw=text(passport), passport_key=passport_key(passport),
             loyalty_cards_json=compact(cards or []))
    p["birth_date_iso"] = iso_birth(birth, birth_fmt, issues)
    if p["passport_raw"].strip().casefold() not in MISSING_ID and not p["passport_key"]:
        issues.append("passport_key_not_parsed")
    if any(not c["card_key"] for c in (cards or [])):
        issues.append("loyalty_key_not_parsed")
    p["issues_json"] = compact(issues)
    return p


def add_issue(row, issue):
    row["issues_json"] = compact(json.loads(row.get("issues_json", "[]")) + [issue])


def observation(p, **flight):
    row = {**p, "record_kind": "user_flight", **flight}
    for key in ("departure_airport", "arrival_airport"):
        if row.get(key) and not re.fullmatch(r"[A-Z]{3}", row[key]):
            add_issue(row, f"{key}_nonstandard")
    return row


class UniqueDict(dict):
    def __setitem__(self, key, value):
        if key in self:
            raise ValueError(f"Duplicate JSON key: {key!r}")
        super().__setitem__(key, value)


def json_profiles(path):
    # Ready-made streaming parser: one profile in memory, not the 568 MB document.
    with path.open("rb") as stream:
        if stream.read(3) != b"\xef\xbb\xbf":
            stream.seek(0)
        def checked_events():
            root_keys = 0
            for prefix,event,value in ijson.parse(stream):
                if prefix == "" and event == "map_key":
                    root_keys += 1
                    if value != "Forum Profiles" or root_keys != 1:
                        raise ValueError("Unexpected JSON root field; refusing to omit it")
                if prefix == "Forum Profiles" and event not in ("start_array","end_array"):
                    raise ValueError("Forum Profiles must be an array")
                yield prefix,event,value
            if root_keys != 1:
                raise ValueError("Expected Forum Profiles array at JSON root")
        yield from ijson.items(checked_events(), "Forum Profiles.item", map_type=UniqueDict)


def boarding(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, delimiter=";")
        header = next(reader)
        if header != BOARDING_COLUMNS:
            raise ValueError("Boarding CSV header differs from documented schema")
        for values in reader:
            if not values:
                continue  # Formatting blank line is not a passenger record.
            if len(values) != len(header):
                raise ValueError(f"Boarding CSV: wrong field count at line {reader.line_num}")
            raw = dict(zip(header, values))
            p = person(raw["PassengerFirstName"], raw["PassengerSecondName"], raw["PassengerLastName"],
                       birth=raw["PassengerBirthDate"], birth_fmt="%m/%d/%Y", sex=raw["PassengerSex"], passport=raw["PassengerDocument"])
            row = observation(p, flight_date=raw["FlightDate"], flight_time=raw["FlightTime"],
                              flight_number=raw["FlightNumber"], arrival_city=raw["Destination"],
                              booking_code=raw["BookingCode"], ticket_number=raw["TicketNumber"],
                              baggage_status=raw["Baggage"], codeshare_raw=raw["CodeShare"])
            yield str(reader.line_num), "json", compact(raw), [("passenger", "$", "none", row)]


def sirena(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        header = stream.readline().rstrip("\r\n")
        if len(header) != 336 or any(header[a:b].strip() != key for key,a,b in SIRENA_COLUMNS):
            raise ValueError("Sirena fixed-width header differs from documented schema")
        for line_number, line in enumerate(stream, 2):
            line = line.rstrip("\r\n")
            if not line.strip():
                continue
            if len(line) != 336:
                raise ValueError(f"Sirena line {line_number}: expected 336 characters, got {len(line)}")
            raw = {key: line[a:b].strip() for key,a,b in SIRENA_COLUMNS}
            tokens = raw["PaxName"].split()
            last, first = (tokens + ["", ""])[:2]
            cards = [card(m[1],m[2]) for m in re.finditer(r"FF#([A-Z0-9]+)\s+(\d+)", raw["PaxAdditionalInfo"])]
            p = person(first, " ".join(tokens[2:]), last, raw["PaxName"], raw["PaxBirthDate"], passport=raw["TravelDoc"], cards=cards)
            if "FF#" in raw["PaxAdditionalInfo"] and not cards:
                add_issue(p,"loyalty_key_not_parsed")
            row = observation(p, flight_date=raw["DepartDate"], flight_time=raw["DepartTime"],
                              arrival_date=raw["ArrivalDate"], arrival_time=raw["ArrivalTime"],
                              flight_number=raw["Flight"], departure_airport=raw["From"], arrival_airport=raw["Dest"],
                              booking_code=raw["Code"], ticket_number=raw["e-Ticket"], codeshare_raw=raw["CodeSh"],
                              travel_class=raw["TrvCls"], fare=raw["Fare"], baggage_allowance=raw["Baggage"],
                              seat=raw["Seat"], meal=raw["Meal"], agent=raw["AgentInfo"])
            # Entire fixed-width line preserves padding and auxiliary fields too.
            yield str(line_number), "fixed_width_text", line, [("passenger","$","passenger",row)]


class NoDTDReader:
    """Reject DTD/entity declarations, including declarations split across reads."""
    def __init__(self, stream):
        self.stream, self.tail = stream, b""

    def read(self, size=-1):
        chunk = self.stream.read(size)
        probe = (self.tail + chunk).upper()
        if b"<!DOCTYPE" in probe or b"<!ENTITY" in probe:
            raise ValueError("XML DTD/entities are not supported")
        self.tail = probe[-16:]
        return chunk


def pointz(path):
    with path.open("rb") as stream:
        parser = ET.iterparse(NoDTDReader(stream), events=("start","end"))
        root = None
        for event, element in parser:
            if root is None:
                root = element
                if root.tag != "PointzAggregatorUsers" or root.attrib:
                    raise ValueError("Unexpected XML root or root attributes")
            if event != "end":
                continue
            if element is not root and element.tag != "user":
                continue
            if element is root:
                # No silent dropping of additional top-level nodes.
                if len(root):
                    raise ValueError("Unexpected XML top-level element")
                break
            name = element.find("name")
            card_nodes = element.findall("./cards/card")
            cards = [combined_card(c.get("number"), bonus_program_raw=c.findtext("bonusprogramm")) for c in card_nodes]
            p = person(first=name.get("first","") if name is not None else "", last=name.get("last","") if name is not None else "", cards=cards)
            uid = element.get("uid","")
            rows = []
            for ci,c in enumerate(card_nodes,1):
                activities = c.findall("./activities/activity")
                for ai,a in enumerate(activities,1):
                    specific = {**p,"loyalty_cards_json":compact([cards[ci-1]])}
                    row = observation(specific, flight_date=a.findtext("Date",""), flight_number=a.findtext("Code",""),
                                      departure_airport=a.findtext("Departure",""), arrival_airport=a.findtext("Arrival",""), fare=a.findtext("Fare",""))
                    if a.get("type") != "Flight":
                        row["record_kind"] = "nonflight_activity"
                        add_issue(row,"activity_type_not_flight")
                    rows.append(("user",f"cards/card[{ci}]/activities/activity[{ai}]","activity_card",row))
            if not rows:
                rows.append(("user","$","profile",{**p,"record_kind":"profile_without_flight"}))
            # XML text includes unknown fields and attributes, not just mapped fields.
            yield f"user[{uid}]", "xml", ET.tostring(element,encoding="unicode"), rows
            root.remove(element)
            element.clear()


def forum(path):
    for index, raw in enumerate(json_profiles(path),1):
        real = raw.get("Real Name") or {}
        cards = [card(c.get("programm"),c.get("Number"), status_raw=c.get("Status")) for c in raw.get("Loyality Programm",[])]
        p = person(first=real.get("First Name"), last=real.get("Last Name"), sex=raw.get("Sex"), cards=cards)
        # Current source has only null passports; preserve and flag future shapes.
        if any(d.get("Passports") is not None for d in raw.get("Travel Documents",[])):
            add_issue(p,"passport_structure_requires_mapping")
        rows = []
        for fi,f in enumerate(raw.get("Registered Flights",[]),1):
            dep, arr = f.get("Departure") or {}, f.get("Arrival") or {}
            row = observation(p, flight_date=f.get("Date",""), flight_number=f.get("Flight",""),
                              codeshare_raw=compact(f["Codeshare"]) if "Codeshare" in f else "",
                              departure_airport=dep.get("Airport",""), departure_city=dep.get("City",""), departure_country=dep.get("Country",""),
                              arrival_airport=arr.get("Airport",""), arrival_city=arr.get("City",""), arrival_country=arr.get("Country",""))
            rows.append(("profile",f"Registered Flights[{fi}]","profile_not_flight",row))
        if not rows:
            rows.append(("profile","$","profile_not_flight",{**p,"record_kind":"profile_without_flight"}))
        yield f"Forum Profiles[{index}]", "json", compact(raw), rows


class UniqueSafeLoader(getattr(yaml, "CSafeLoader", yaml.SafeLoader)):
    pass


def yaml_mapping(loader, node):
    pairs = loader.construct_pairs(node, deep=True)
    result = {}
    for key,value in pairs:
        if key in result:
            raise ValueError(f"Duplicate YAML key: {key!r}")
        result[key] = value
    return result


UniqueSafeLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, yaml_mapping)


def yaml_days(path):
    # This export has dates at indentation zero. Parse one day's YAML at a time.
    lines = []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        for line in stream:
            if line.strip() and not line[0].isspace():
                if not re.fullmatch(r"'\d{4}-\d{2}-\d{2}':\s*", line):
                    raise ValueError("Unexpected YAML root: expected a quoted date")
                if lines and any(part.strip() for part in lines):
                    yield "".join(lines)
                lines = []
            lines.append(line)
        if lines:
            yield "".join(lines)


def skyteam(path):
    for raw_day in yaml_days(path):
        day = yaml.load(raw_day, Loader=UniqueSafeLoader)
        date,flights = next(iter(day.items()))
        rows = []
        for flight,fields in flights.items():
            common = dict(flight_date=date,flight_number=flight,departure_airport=fields["FROM"],
                          arrival_airport=fields["TO"],flight_status=fields["STATUS"])
            for ci,(number,details) in enumerate(fields["FF"].items(),1):
                p = person(cards=[combined_card(number)])
                rows.append((f"{flight}:ff{ci}",f"{date}/{flight}/FF[{ci}]","flight_participant",
                             observation(p,**common,travel_class=details["CLASS"],fare=details["FARE"])))
            if not fields["FF"]:
                rows.append((f"{flight}:anonymous",f"{date}/{flight}","none",
                             observation(person(),**common,record_kind="flight_without_user")))
        # Keep the entire day, including unmapped attributes, as original YAML text.
        yield date,"yaml",raw_day,rows


def excel_cell_text(cell):
    value = cell.value
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    # Excel may store an identifier as an integer with a leading-zero display mask.
    if isinstance(value, int) and re.fullmatch(r"0+", cell.number_format or ""):
        return str(value).zfill(len(cell.number_format))
    return text(value)


def boardingpass(path):
    with zipfile.ZipFile(path) as archive:
        members = [m for m in archive.infolist() if not m.is_dir()]
        if not members or any(Path(m.filename).suffix.lower() != ".xlsx" for m in members):
            raise ValueError("Expected only XLSX workbooks in boarding-pass archive")
        for member in members:
            print(f"[boardingpass] Reading {member.filename}", flush=True)
            # One workbook in memory; no extraction to disk and no path handling.
            book = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=False)
            try:
                for sheet in book.worksheets:
                    cells = {cell.coordinate: cell for row in sheet.iter_rows() for cell in row if cell.value is not None}
                    get = lambda key: excel_cell_text(cells[key]) if key in cells else ""
                    labels = {"A1":"BOARDING PASS", "F1":"SEQUENCE:", "A7":"GATE", "G11":"SEAT", "A13":"PNR:", "D13":"E-TICKET"}
                    if any(get(key) != label for key,label in labels.items()):
                        raise ValueError(f"Unexpected boarding-pass layout: {member.filename}/{sheet.title}")
                    loyalty = get("F3")
                    cards = [combined_card(loyalty)] if loyalty.strip().casefold() not in MISSING_ID else []
                    p = person(full=get("B3"), cards=cards)
                    add_issue(p, "name_order_not_inferred")
                    if not cards:
                        add_issue(p, "identity_key_missing")
                    row = observation(p, flight_date=get("A9"), flight_time=get("C9"), flight_number=get("A5"),
                                      departure_airport=get("D7"), departure_city=get("D5"), arrival_airport=get("H7"), arrival_city=get("H5"),
                                      codeshare_raw=get("E9"), travel_class=get("H3"), seat=get("H11"), gate=get("B7"),
                                      boarding_sequence=get("H1"), honorific=get("A3"), booking_code=get("B13"), ticket_number=get("E13"))
                    raw_cells = {}
                    for coordinate,cell in cells.items():
                        value = cell.value
                        if isinstance(value, (datetime, date, datetime_time)):
                            value = {"python_type":type(value).__name__, "iso":value.isoformat()}
                        raw_cells[coordinate] = {"value":value, "excel_type":cell.data_type, "number_format":cell.number_format}
                    payload = compact({"workbook":member.filename,"sheet":sheet.title,"cells":raw_cells})
                    yield f"{member.filename}/{sheet.title}", "xlsx_cells_json", payload, [("passenger",sheet.title,"passenger",row)]
            finally:
                book.close()


CONVERTERS = {"boarding":boarding,"sirena":sirena,"pointz":pointz,"forum":forum,"skyteam":skyteam,"boardingpass":boardingpass}


def csv_writer(stack,path,columns):
    stream = stack.enter_context(path.open("w",encoding="utf-8-sig",newline=""))
    writer = csv.DictWriter(stream,fieldnames=columns,delimiter=";",extrasaction="ignore",quoting=csv.QUOTE_MINIMAL)
    writer.writeheader()
    return writer


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda:stream.read(4*1024*1024),b""):
            digest.update(chunk)
    return digest.hexdigest()


def run(input_dir: Path, output_root: Path, selected=None, limit_records=0) -> Path:
    selected = list(selected or SOURCES)
    if limit_records<0:
        raise ValueError("limit_records must be nonnegative")
    paths = {s:input_dir/SOURCES[s] for s in selected}
    for path in paths.values():
        if not path.is_file():
            raise FileNotFoundError(path)
    output_root.mkdir(parents=True,exist_ok=True)
    output = output_root / (datetime.now().strftime("run_%Y%m%d_%H%M%S_%f") + ("_SAMPLE" if limit_records else ""))
    output.mkdir()
    (output/"INCOMPLETE.txt").write_text("Run incomplete. Do not merge these outputs.\n",encoding="utf-8")
    manifest = []
    try:
        with ExitStack() as stack:
            for source,path in paths.items():
                before = path.stat()
                started = time.monotonic()
                counts = Counter()
                print(f"[{source}] Starting {path.name}",flush=True)
                with ExitStack() as files:
                    main = csv_writer(files,output/f"{source}.csv",DATA_COLUMNS)
                    records = csv_writer(files,output/f"{source}.records.csv",RECORD_COLUMNS)
                    links = csv_writer(files,output/f"{source}.links.csv",LINK_COLUMNS)
                    generator = CONVERTERS[source](path)
                    try:
                        for index, (position,fmt,payload,rows) in enumerate(generator,1):
                            rid = f"{source}:r{index:09d}"
                            records.writerow(dict(source_record_id=rid,source_file=path.name,source_position=position,payload_format=fmt,payload=payload))
                            counts["source_records"]+=1
                            for ri,(local,locator,scope,row) in enumerate(rows,1):
                                oid = f"{rid}:o{ri:04d}"
                                main.writerow({**row,"observation_id":oid})
                                links.writerow(dict(observation_id=oid,source_record_id=rid,source_subject_id=f"{rid}:{local}",source_path=locator,loyalty_scope=scope))
                                counts["observations"]+=1
                                counts["rows_with_issues"]+=row.get("issues_json","[]")!="[]"
                            if index % 10000 == 0:
                                print(f"[{source}] {index:,} source records; {counts['observations']:,} observations",flush=True)
                            if limit_records and index>=limit_records:
                                break
                    finally:
                        generator.close()
                after = path.stat()
                if (before.st_size,before.st_mtime_ns)!=(after.st_size,after.st_mtime_ns):
                    raise RuntimeError(f"Input changed during conversion: {path.name}")
                manifest.append(dict(source=source,source_file=str(path.resolve()),input_bytes=before.st_size,
                                     input_sha256=sha256(path) if not limit_records else "not_computed_in_sample_mode",
                                     sample_limit_records=limit_records,**counts,seconds=round(time.monotonic()-started,2),converter_version=VERSION))
                print(f"[{source}] Done: {counts['observations']:,} observations",flush=True)
            mw = csv_writer(stack,output/"manifest.csv",[
                "source","source_file","input_bytes","input_sha256","sample_limit_records","source_records",
                "observations","rows_with_issues","seconds","converter_version"])
            mw.writerows(manifest)
            status = csv_writer(stack,output/"SUCCESS.csv",["converter_version","mode","sources","global_users_assigned"])
            status.writerow(dict(converter_version=VERSION,mode="sample" if limit_records else "full",sources=len(selected),global_users_assigned="false"))
        (output/"INCOMPLETE.txt").unlink()
        return output
    except BaseException:
        raise


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir",type=Path,default=INPUT_DIR)
    parser.add_argument("--output-dir",type=Path,default=OUTPUT_ROOT)
    parser.add_argument("--sources",nargs="+",choices=list(SOURCES),default=list(SOURCES))
    parser.add_argument("--limit-records",type=int,default=0,help="Per-source record limit for a quick sample; 0 = all records")
    args = parser.parse_args()
    try:
        result = run(args.input_dir,args.output_dir,args.sources,args.limit_records)
    except (Exception,KeyboardInterrupt) as exc:
        print(f"FAILED: {exc}\nOriginals unchanged. Any INCOMPLETE output must not be merged.",file=sys.stderr)
        return 1
    print(f"SUCCESS: {result}",flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
